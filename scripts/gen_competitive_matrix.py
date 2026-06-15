#!/usr/bin/env python3
"""Generate the Echelon vs CrewAI vs Hatz AI competitive feature matrix workbook.

Produces docs/research/2026-06-15-echelon-crewai-hatz-comparison.xlsx
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---- styling helpers -------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
CAT_FILL = PatternFill("solid", fgColor="374151")
CAT_FONT = Font(color="FFFFFF", bold=True, size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# alignment-tag colours
TAG_FILL = {
    "STRENGTH": PatternFill("solid", fgColor="D1FAE5"),   # green
    "UNIQUE":   PatternFill("solid", fgColor="DBEAFE"),   # blue
    "PARTIAL":  PatternFill("solid", fgColor="FEF3C7"),   # amber
    "GAP":      PatternFill("solid", fgColor="FEE2E2"),   # red
    "DIFFERENT":PatternFill("solid", fgColor="EDE9FE"),   # purple
}
EFFORT_FILL = {
    "S":   PatternFill("solid", fgColor="D1FAE5"),
    "M":   PatternFill("solid", fgColor="DCFCE7"),
    "L":   PatternFill("solid", fgColor="FEF3C7"),
    "XL":  PatternFill("solid", fgColor="FED7AA"),
    "XXL": PatternFill("solid", fgColor="FECACA"),
    "N/A": PatternFill("solid", fgColor="F3F4F6"),
}

wb = openpyxl.Workbook()

# ===========================================================================
# SHEET 1 — README / Legend
# ===========================================================================
ws = wb.active
ws.title = "README"
ws.sheet_view.showGridLines = False
readme = [
    ("Echelon vs CrewAI vs Hatz AI — Competitive Feature Matrix", "title"),
    ("Prepared 2026-06-15  ·  Echelon baseline = main @ v0.1.0 + Plan 11 host work", "sub"),
    ("", ""),
    ("HOW TO READ THIS WORKBOOK", "h"),
    ("Tab 'Feature Comparison' is the core matrix: every row is a capability, scored for each "
     "of the three products, tagged for alignment, and given an effort estimate to build/refine "
     "into Echelon. Tab 'Build Roadmap' is the prioritised shortlist of what to actually pursue. "
     "Tab 'Platform Profiles' is a one-screen summary of each product.", "p"),
    ("", ""),
    ("ALIGNMENT TAG KEY", "h"),
    ("STRENGTH", "Echelon already has this and is at parity or ahead of the others."),
    ("UNIQUE", "Echelon-only — neither CrewAI nor Hatz has a real equivalent. A moat to protect."),
    ("PARTIAL", "Echelon has it designed/scaffolded but it is STUBBED or incomplete in v0.1."),
    ("GAP", "Missing in Echelon; present in one or both competitors. Candidate to build."),
    ("DIFFERENT", "All have a take, but the approaches diverge intentionally — not a simple gap."),
    ("", ""),
    ("EFFORT KEY (to design + build + reverse-engineer into Echelon, 1 engineer)", "h"),
    ("S", "~1-3 days. Config, wiring, or a thin adapter over something that exists."),
    ("M", "~1-2 weeks. A self-contained feature/service with tests."),
    ("L", "~3-6 weeks. A subsystem touching several layers (renderer + main + MCP)."),
    ("XL", "~2-3 months. A new architectural pillar (e.g. observability stack, eval harness)."),
    ("XXL", "~1 quarter+. Strategic re-platforming (e.g. multi-tenant SaaS, on-prem control plane)."),
    ("N/A", "Not recommended for Echelon, or a deliberate non-goal given its local-first thesis."),
    ("", ""),
    ("KEY CONTEXT — three very different products", "h"),
    ("CrewAI", "Open-source Python multi-agent FRAMEWORK (MIT) + commercial AMP control plane. "
     "Developer-facing. Code/YAML authored teams. ~$18M raised, 60%+ Fortune-500 reach claimed."),
    ("Hatz AI", "Closed-source multi-tenant white-label SaaS sold THROUGH MSPs to SMBs. "
     "No public product code (only a beta TS SDK). Business-user facing. SOC2 Type I/II, SSO/SCIM."),
    ("Echelon", "Local-first Electron + Next.js DESKTOP app, BYO API keys, no cloud/subscription. "
     "Themed character agent teams, 4-model counselor, review gates, seasons. Single-user today."),
    ("", ""),
    ("HONESTY NOTE ON ECHELON v0.1", "h"),
    ("Per docs/CODEBASE_MAP.md several flagship pieces are scaffolded but not yet functional: "
     "counselor-service returns MOCK verdicts, the 7 review gates are STUBS, 'kb promote' is "
     "deferred to v0.5, and there is no auth/RBAC/observability/deployment layer. The matrix "
     "scores current reality (PARTIAL) but notes the intended design.", "p"),
]
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 110
r = 1
for left, right in readme:
    c1 = ws.cell(row=r, column=1, value=left)
    c2 = ws.cell(row=r, column=2, value=right if right not in ("title","sub","h","p","") else "")
    if right == "title":
        c1.font = Font(bold=True, size=16, color="1F2937")
    elif right == "sub":
        c1.font = Font(italic=True, size=10, color="6B7280")
    elif right == "h":
        c1.font = Font(bold=True, size=12, color="1F2937")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    elif right == "p":
        c1.font = Font(size=10, color="374151")
        c1.alignment = WRAP
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.row_dimensions[r].height = 45
    else:
        c1.font = Font(bold=True, size=10)
        c2.font = Font(size=10, color="374151")
        c2.alignment = WRAP
        if left in TAG_FILL:
            c1.fill = TAG_FILL[left]
        if left in EFFORT_FILL:
            c1.fill = EFFORT_FILL[left]
        c1.alignment = CENTER
    r += 1

# ===========================================================================
# SHEET 2 — Feature Comparison (the matrix)
# ===========================================================================
ws = wb.create_sheet("Feature Comparison")
cols = ["#", "Category", "Capability", "Echelon (v0.1)", "CrewAI", "Hatz AI",
        "Alignment", "Effort", "Recommendation / Notes"]
widths = [4, 20, 26, 40, 40, 38, 12, 8, 52]
for i, (c, w) in enumerate(zip(cols, widths), start=1):
    cell = ws.cell(row=1, column=i, value=c)
    cell.fill = HEADER_FILL; cell.font = HEADER_FONT
    cell.alignment = CENTER; cell.border = BORDER
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"

# rows: (category, capability, echelon, crewai, hatz, tag, effort, note)
# category left blank to repeat = visual grouping handled below
ROWS = [
 ("Orchestration & Execution", "Multi-agent orchestration paradigm",
  "Seasons of themed characters cast from 43 archetypes; orchestrator MCP + super-agent delegates work.",
  "Two models: 'Crews' (autonomous role-based teams) and 'Flows' (deterministic event-driven workflows). Industry-leading clarity.",
  "Workflows + agents + AI apps built no-code; multi-department orchestration, voice agents.",
  "STRENGTH", "—",
  "Echelon's roster/season model is strong & distinctive. Borrow CrewAI's explicit Crews-vs-Flows split as a mental model."),

 ("Orchestration & Execution", "Deterministic workflow engine (Flows-equivalent)",
  "No first-class deterministic workflow DSL; automations + kanban auto-spawn are ad-hoc.",
  "Flows: @start/@listen/@router/@persist, Pydantic-typed state, or_/and_ combinators, resume & fork, plot() viz.",
  "Step-based workflows; re-run from any step; post-call workflow chains.",
  "GAP", "L",
  "HIGH VALUE. A typed, inspectable workflow layer above seasons would make orchestration deterministic & debuggable. Model on CrewAI Flows."),

 ("Orchestration & Execution", "Hierarchical / manager delegation",
  "Super-agent orchestrator + convener per theme; delegate_task primitive.",
  "Hierarchical process with manager_llm — but this is CrewAI's MOST-criticized feature (poor coordination, token waste, delegation bugs).",
  "Multi-department orchestration (managed).",
  "STRENGTH", "—",
  "Echelon's convener/counselor escalation is arguably cleaner than CrewAI's flaky manager. Avoid CrewAI's mistakes here."),

 ("Orchestration & Execution", "Parallel execution / isolation",
  "Git worktree-per-agent; multi-season isolation; 10+ concurrent agents.",
  "async_execution on tasks; no built-in worktree/workspace isolation.",
  "Tenant/org/user logical isolation (SaaS).",
  "UNIQUE", "—",
  "Worktree-per-agent isolation is an Echelon strength neither competitor matches. Protect & market it."),

 ("Orchestration & Execution", "Event-driven triggers / automations",
  "Automation engine polls GitHub, JIRA; stubbed Pipedrive/Twitter/RSS; spawns disposable agents.",
  "Flows event model; webhook streaming of events (AMP).",
  "Zapier / Make / n8n bridges; native triggers across 40+ integrations.",
  "STRENGTH", "M",
  "Echelon already has this. Add Zapier/Make/n8n bridges (Hatz pattern) to widen trigger surface cheaply."),

 ("Orchestration & Execution", "Scheduling (cron)",
  "Scheduler with cron CRUD + logs + watch (mature).",
  "AMP cron scheduling in control plane.",
  "Scheduled/recurring workflows.",
  "STRENGTH", "—", "At parity. No action."),

 ("Orchestration & Execution", "Task board / kanban",
  "@dnd-kit kanban with agent auto-sync & auto-spawn on 'planned'.",
  "No native board (tasks are code).",
  "Workflow/queue views.",
  "UNIQUE", "—", "Echelon-unique visual task→agent binding. Keep; add debounce (known bug)."),

 ("Agent Definition & Personas", "Declarative agent config",
  "YAML archetypes + 5-7 file soul packages (SOUL/AGENTS/HEARTBEAT/MEMORY/persona); schema-validated.",
  "YAML agents.yaml/tasks.yaml + @CrewBase wiring; {var} interpolation.",
  "No-code app builder; prompt/templates.",
  "STRENGTH", "—",
  "Echelon's soul-package model is richer than CrewAI's role/goal/backstory. Strength."),

 ("Agent Definition & Personas", "Persona / character identity",
  "Themed recognizable characters (TBBT, Young Sheldon) with voice, mannerisms, guardrails.",
  "Flat role/goal/backstory strings.",
  "Generic branded assistants.",
  "UNIQUE", "—", "Signature Echelon differentiator. Neither competitor has relatable character UX."),

 ("Agent Definition & Personas", "Role/archetype library",
  "43 archetypes, tiered (medium/large/enterprise) for auto-roster sizing.",
  "Users define roles ad-hoc; blueprints in AMP Discovery.",
  "Pre-built workflow/app templates.",
  "STRENGTH", "—", "Strong. Consider publishing as shareable blueprints (see Templates row)."),

 ("Agent Definition & Personas", "PRD → roster composition",
  "roster-composer parses a PRD and selects/sizes the team automatically.",
  "Crew Studio: NL prompt → crew in <60s; Discovery suggests blueprints.",
  "NL app builder.",
  "STRENGTH", "M",
  "Echelon's PRD→roster is unique & valuable. Add CrewAI-style NL conversational refinement on top."),

 ("Memory & Knowledge", "Cross-session memory",
  "MEMORY.seed.md drifts at runtime; mempalace KB; auto-memory MEMORY.md.",
  "Unified Memory class: hierarchical scopes, recency/semantic/importance weighting, consolidation dedup.",
  "Memory-based RAG; auto-vectorize over-context files.",
  "DIFFERENT", "L",
  "Borrow CrewAI's weighted/scoped memory math (recency_half_life, importance) to upgrade Echelon's KB ranking."),

 ("Memory & Knowledge", "Knowledge base / RAG",
  "mempalace-backed KB (LOCAL solo mode v0.1); kb query/write; promote-skill deferred to v0.5.",
  "Knowledge sources (PDF/CSV/Excel/JSON/web), ChromaDB/Qdrant, query rewriting, score thresholds.",
  "Pinecone + Snowflake Cortex; 60+ file types; audio; auto-vectorization.",
  "PARTIAL", "L",
  "Echelon KB is local-only & partly stubbed. Add multi-format ingestion (CrewAI source classes) + query rewriting."),

 ("Memory & Knowledge", "Knowledge flywheel / accumulation",
  "Captures learnings on merge; queries prior art before tasks; advisory/counselor verdicts shared globally.",
  "No cross-project institutional-memory flywheel by default.",
  "Community workflow sharing (manual).",
  "UNIQUE", "—", "Echelon's compounding-memory thesis is a genuine moat. Make it real (it's partly stubbed)."),

 ("Memory & Knowledge", "Embedder / vector-store options",
  "mempalace built-in index; single backend.",
  "11+ embedders (OpenAI/Vertex/Cohere/Voyage/Bedrock/HF/Jina/WatsonX), pluggable StorageBackend.",
  "Pinecone, Snowflake Cortex.",
  "GAP", "M",
  "Add a pluggable embedder/vector-store abstraction so power users can bring their own."),

 ("Memory & Knowledge", "Document handling breadth",
  "Vault (SQLite) doc store; attachFile; folders/search.",
  "PDF/CSV/Excel/JSON/web via Docling.",
  "60+ file types incl. legacy Office, archives, code, AUDIO (mp3/wav/...).",
  "GAP", "M",
  "Hatz's 60+ types + audio is the bar. Expand vault ingestion + add an audio/transcription path."),

 ("Multi-Model / LLM", "Multi-provider LLM support",
  "CLI provider abstraction: Claude, Codex, Gemini, OpenCode, Pi.",
  "Native: OpenAI/Anthropic/Gemini/Azure/Bedrock/Snowflake; +LiteLLM for the rest.",
  "'Mido' gateway: OpenAI/Anthropic/Amazon/Google/DeepSeek/Grok.",
  "STRENGTH", "S",
  "Good coverage. Consider a LiteLLM-style fallback router to instantly widen support."),

 ("Multi-Model / LLM", "Per-agent model selection",
  "Provider per agent via CLI provider; not surfaced per-character cleanly.",
  "Per-agent llm + separate function_calling_llm.",
  "Per-task model picker with cost multipliers shown.",
  "DIFFERENT", "S",
  "Surface a per-character model picker in the UI; show relative cost (Hatz multipliers idea)."),

 ("Multi-Model / LLM", "Multi-model consensus council",
  "Counselor: 4 models (Gemini/GPT-5/Opus/Grok), 4 placement types, consensus rules — but STUBBED (mock verdicts).",
  "Single-model per task; no consensus council.",
  "Single model per task.",
  "UNIQUE", "M",
  "FLAGSHIP differentiator nobody else has — but it returns mock data today. Wiring real fan-out is the #1 unlock."),

 ("Multi-Model / LLM", "LLM gateway w/ cost multipliers",
  "No gateway; keys per provider.",
  "LiteLLM routing; usage_metrics.",
  "Mido gateway brokers providers, shows credit multipliers, insulates from vendor policy.",
  "GAP", "L",
  "A lightweight internal gateway (routing + per-call cost accounting) would power budgets & analytics."),

 ("Multi-Model / LLM", "Local model support",
  "Tasmania local LLM client; opencode/pi providers.",
  "Ollama, NVIDIA NIM.",
  "Not emphasized (cloud SaaS).",
  "STRENGTH", "—", "Echelon's local-LLM path fits its local-first thesis. Strength."),

 ("Quality, Review & Guardrails", "Review gates / quality pipeline",
  "7 gates (architecture/code/QA/security/adversarial/UI/refinement) — all STUBBED in v0.1.",
  "No built-in multi-gate review pipeline.",
  "No code-review pipeline (business AI).",
  "PARTIAL", "L",
  "Designed & unique, but stubs. Implementing the gate runner for real is a top-3 priority."),

 ("Quality, Review & Guardrails", "Hallucination / faithfulness guardrail",
  "None (security gate stubbed).",
  "Hallucination Guardrail: 0-10 faithfulness score vs reference, auto-runs before completion, retries.",
  "Not advertised.",
  "GAP", "M",
  "Borrow directly. A faithfulness scorer fits Echelon's quality-gate framing perfectly."),

 ("Quality, Review & Guardrails", "Human-in-the-loop approval",
  "User-handler merge approval; override CLI; review-gate updated push.",
  "human_input on tasks; Flow @human_feedback() pause/resume.",
  "Admin approval flows.",
  "STRENGTH", "S", "At/above parity. Add Flow-style typed pause/resume if Flows layer is built."),

 ("Quality, Review & Guardrails", "Output guardrails (function/LLM)",
  "Soul-package hard guardrails (prompt-level); no programmatic output validators.",
  "Function-based + LLM-based guardrails, chainable, guardrail_max_retries.",
  "Role/integration controls.",
  "GAP", "M",
  "Add programmatic + LLM output validators on agent/task results (CrewAI signature pattern)."),

 ("Quality, Review & Guardrails", "Structured outputs",
  "YAML schemas for artifacts; no runtime Pydantic-style typed outputs.",
  "output_json / output_pydantic via Pydantic + instructor.",
  "Structured app outputs.",
  "GAP", "S", "Add typed/JSON-schema output contracts for agent results. Cheap, high leverage for reliability."),

 ("Developer Experience", "CLI scaffolding",
  "9 CLI families (season/character/scope/kb/counselor/...).",
  "crewai create/run/deploy/login/org/config — rich.",
  "No public dev CLI (SaaS).",
  "STRENGTH", "—", "At parity. Strong CLI surface already."),

 ("Developer Experience", "No-code / visual builder",
  "Rich UIs (canvas, kanban, terminals) but no visual TEAM/flow builder.",
  "Crew Studio: NL copilot + drag-drop canvas + Python export.",
  "Drag-and-drop AI App Builder (no-code).",
  "GAP", "XL",
  "Both competitors lean hard on no-code builders. A visual season/flow builder broadens Echelon's audience beyond devs."),

 ("Developer Experience", "Testing / eval harness",
  "Vitest + Bun suites for the app; no AGENT-output eval harness.",
  "crewai test: runs crew N times, LLM scores per-task/agent.",
  "Not exposed.",
  "GAP", "L",
  "An LLM-scored agent-eval loop (CrewAI pattern) would quantify season quality over time. Pairs with the flywheel."),

 ("Developer Experience", "Training (human feedback)",
  "None.",
  "crewai train: human-feedback loop persisted to .pkl (reliability complaints noted).",
  "None.",
  "GAP", "L",
  "Lower priority — CrewAI's own impl is buggy. Echelon's KB flywheel is a better-aligned learning mechanism."),

 ("Developer Experience", "Replay / resume",
  "rerun CLI family; season state on disk.",
  "crewai replay -t task_id; Flow persist/resume/fork.",
  "Re-run workflow from any step.",
  "PARTIAL", "M", "Formalize replay-from-step + resume across all agent runs (Flow persist pattern)."),

 ("Developer Experience", "Templates / blueprints marketplace",
  "Archetypes/themes are extensible YAML; skills.sh marketplace; no team blueprints exchange.",
  "AMP blueprints from billions of deployments; Tool Repository.",
  "Community workflow marketplace; MSP-published templates; AI-in-a-Box kit.",
  "GAP", "L",
  "A shareable season/blueprint registry creates network effects. Both competitors invest here."),

 ("Developer Experience", "API / SDK",
  "HTTP API on :31415 (bearer auth); REST routes for agent control.",
  "REST Kickoff API + webhook streaming.",
  "Beta TypeScript SDK (models/apps/completions/files).",
  "STRENGTH", "S", "Have it. Publish a typed client SDK + OpenAPI for external integrators."),

 ("Developer Experience", "MCP support",
  "7 MCP servers (orchestrator/kanban/vault/telegram/socialdata/x/world); custom MCP config.",
  "First-class MCP (managed + manual), thousands of servers.",
  "Custom MCP servers w/ role-based access; public remote-mcp-server.",
  "STRENGTH", "—", "Echelon is MCP-native already. Strength; add role-based MCP access (Hatz)."),

 ("Observability & Monitoring", "Execution tracing",
  "Per-agent xterm output; statusline; no structured traces.",
  "AMP traces (crewai traces); OTEL backbone; +Langfuse/Phoenix/Weave/MLflow integrations.",
  "Audit trails / logging.",
  "GAP", "XL",
  "Biggest enterprise gap. Add OpenTelemetry-based structured tracing of agent/tool/LLM spans. Foundation for everything below."),

 ("Observability & Monitoring", "Usage / cost analytics",
  "Token/cost tracking + activity patterns (usage tracking shipped).",
  "Usage dashboards; usage_metrics; outcome tracking.",
  "Credit metering, per-client usage dashboards, model multipliers.",
  "PARTIAL", "M",
  "Have basics. Build per-season/per-character cost dashboards; pairs with a gateway for accurate accounting."),

 ("Observability & Monitoring", "Real-time monitoring dashboard",
  "Dashboard component; live agent ticks.",
  "Control-plane real-time metrics/logs/traces.",
  "Admin dashboard + usage views.",
  "PARTIAL", "M", "Extend dashboard to a true ops view (active agents, queue, errors, spend)."),

 ("Observability & Monitoring", "Audit logs",
  "Merges serialized through one authority (audit trail by design); no formal immutable audit log.",
  "Immutable audit trails (AMP).",
  "Audit trails / logging per tenant.",
  "GAP", "M", "Add an append-only audit log of agent actions, approvals, overrides. Needed for any enterprise sale."),

 ("Deployment & Hosting", "Local desktop app",
  "Electron desktop, runs on user machine, BYO keys, no cloud dependency.",
  "Not a desktop app (framework/SaaS).",
  "Not a desktop app (SaaS).",
  "UNIQUE", "—", "Local-first privacy is a real differentiator for security-conscious buyers. Protect it."),

 ("Deployment & Hosting", "Managed cloud / SaaS",
  "None (deliberate local-first thesis).",
  "AMP Cloud, serverless auto-scaling.",
  "Fully managed multi-tenant SaaS.",
  "DIFFERENT", "XXL",
  "Strategic fork-in-road: a hosted Echelon would unlock teams/enterprise but contradicts the current thesis. Decide deliberately."),

 ("Deployment & Hosting", "Self-hosted / on-prem / VPC",
  "Inherently self-hosted (it's local). No server/VPC deployment.",
  "CrewAI Factory: containerized on-prem + private VPC (AWS/Azure/GCP); single-tenant.",
  "AWS US-East only; no on-prem.",
  "DIFFERENT", "XL",
  "If a team/server edition is pursued, containerized self-host (CrewAI Factory model) is the enterprise-friendly path."),

 ("Enterprise: Multi-tenancy & Access", "Multi-tenancy",
  "Single-user; multi-SEASON isolation but not multi-TENANT.",
  "Org/workspace constructs in AMP.",
  "Core strength: tenant/org/user isolation, per-tenant branding & templates.",
  "GAP", "XXL",
  "Requires the SaaS fork. Hatz's tenant/org/user model is the reference if Echelon ever goes multi-tenant."),

 ("Enterprise: Multi-tenancy & Access", "White-label / branding",
  "Theming exists for CHARACTERS, not for reseller branding. (landing/ still Dorothy-branded.)",
  "Limited.",
  "Core: per-partner branded instances; clients never see Hatz.",
  "DIFFERENT", "L",
  "Only relevant under a reseller/agency motion. Echelon's theme engine is a head-start if pursued."),

 ("Enterprise: Multi-tenancy & Access", "RBAC",
  "capabilities/access-matrix.yaml defines capability SCOPES for agents; no human-user RBAC.",
  "Granular RBAC roles in control plane.",
  "RBAC across users + per-role integration/MCP controls.",
  "GAP", "L",
  "Echelon has agent-capability RBAC; lacks human-user RBAC. Needed for any team/enterprise edition."),

 ("Enterprise: Multi-tenancy & Access", "SSO (SAML/OIDC)",
  "None (single-user desktop).",
  "Entra ID, Okta, Auth0.",
  "SAML 2.0 (Okta/Entra/generic), domain verification.",
  "GAP", "L", "Table-stakes for enterprise; only meaningful once there is a hosted/multi-user edition."),

 ("Enterprise: Multi-tenancy & Access", "SCIM provisioning",
  "None.",
  "Not emphasized.",
  "SCIM lifecycle sync from IdP.",
  "GAP", "M", "Follows SSO. Defer until hosted edition exists."),

 ("Enterprise: Multi-tenancy & Access", "MFA",
  "None (local app).",
  "Via customer IdP.",
  "Enforced MFA all tiers + Microsoft OAuth.",
  "GAP", "S", "Trivial relative to hosting; bundled with SSO work."),

 ("Enterprise: Multi-tenancy & Access", "Admin console",
  "15+ section Settings panel (single-user).",
  "Crew Control Plane (governance, RBAC, scheduling, usage).",
  "admin.hatz.ai multi-tenant portal + workspace templates.",
  "DIFFERENT", "L", "A team-admin console is the gateway to enterprise; depends on multi-user model."),

 ("Security & Compliance", "SOC 2",
  "None (no hosted service to certify).",
  "Claimed SOC2 (no public report surfaced — verify).",
  "SOC 2 Type I & II + SOC 3, NDA-gated reports, third-party pentests.",
  "GAP", "XXL",
  "Only applies to a hosted edition. Hatz is the model: actually certified, not just claimed."),

 ("Security & Compliance", "HIPAA / GDPR",
  "Local-first means data never leaves the machine (a privacy story in itself).",
  "Claims HIPAA/FedRAMP High (unverified, no BAA surfaced).",
  "NOT publicly confirmed (open question for regulated buyers).",
  "DIFFERENT", "XL", "Echelon's local-first model sidesteps much of this — a marketing angle, not a gap, for solo use."),

 ("Security & Compliance", "Data isolation / no-train guarantee",
  "Data stays local; no telemetry/phone-home (per cost-model doc).",
  "Telemetry controversy historically; now opt-in + OTEL_SDK_DISABLED.",
  "Inference layer separated from history store; no-train agreements (get in writing).",
  "STRENGTH", "—", "Local-first + no telemetry is a clean privacy story. Make it explicit & audited."),

 ("Security & Compliance", "Sandboxed code execution",
  "Agents run via CLI providers; relies on host; worktree isolation but not a hard sandbox.",
  "External sandboxes (allow_code_execution deprecated); 2026 CVEs around sandbox escape/RCE.",
  "N/A (no code-exec product).",
  "GAP", "L",
  "Learn from CrewAI's CVEs: if Echelon adds code-exec, sandbox properly (no fallback-to-unsafe)."),

 ("Security & Compliance", "Secrets handling",
  "app-settings.json written owner-only (0600); API token file.",
  "Env vars / IdP.",
  "Managed credential store per tenant.",
  "STRENGTH", "S", "Solid baseline. Consider OS keychain integration for keys."),

 ("Integrations", "Business app integrations",
  "Telegram, Slack, GitHub, JIRA, X (post + read).",
  "23+ OAuth (Gmail/Slack/Teams/Jira/Notion/Salesforce/HubSpot/Stripe/Shopify/...).",
  "40+ (M365 deep, Google, ConnectWise, Salesforce, HubSpot, Notion, Linear, ...).",
  "GAP", "L",
  "Echelon trails badly on breadth. Adopt an OAuth-connector framework + the Zapier/Make/n8n bridge to leapfrog cheaply."),

 ("Integrations", "PSA / MSP integrations",
  "None.",
  "None.",
  "ConnectWise, Autotask, HaloPSA, Pax8, ScalePad (channel moat).",
  "DIFFERENT", "N/A", "Only relevant if Echelon ever pursues an MSP channel. Non-goal today."),

 ("Integrations", "Voice / phone agents",
  "None.",
  "Not core.",
  "AI Phone / voice agents with post-call workflows.",
  "GAP", "L", "Novel but off-thesis for a dev-team factory. Park unless product direction shifts."),

 ("Integrations", "Automation bridges (Zapier/Make/n8n)",
  "None (custom webhooks only, with SSRF protection).",
  "Composio integration tool.",
  "Zapier, Make, n8n natively.",
  "GAP", "M", "Cheapest way to 1000s of integrations. High ROI. Build the bridge nodes."),

 ("Business Model / GTM", "Pricing model",
  "Free / BYO API keys; ~$1-5 API cost per season; default $50 counselor budget cap.",
  "Free tier (~50 exec/mo); Pro ~$25; Enterprise custom (unstable public pricing).",
  "Channel-gated, no public price; pooled non-rollover credits; ~$190/mo entry (reported).",
  "DIFFERENT", "—", "Echelon's near-zero-cost local model is a genuine wedge vs per-seat SaaS. Lead with it."),

 ("Business Model / GTM", "Channel / reseller program",
  "None.",
  "Partner motions exist; mainly direct/PLG.",
  "Entire GTM: ~40k MSPs, NFR packages, certifications, partner Discord, AI-in-a-Box.",
  "DIFFERENT", "N/A", "Strategic non-goal unless Echelon productizes for agencies. Note as future option."),

 ("Business Model / GTM", "Community / learning",
  "Slide deck + docs; no academy.",
  "learn.crewai.com, 100k+ certified devs, 50k Discord — strong moat.",
  "Partner enablement, self-paced modules.",
  "GAP", "M", "Content/community compounds adoption. Low-cost, high-leverage; start with docs + examples gallery."),

 ("Echelon-Unique Strengths", "Themed character UX",
  "Recognizable personas with consistent voice/quirks; trust through familiarity.",
  "None.", "None.",
  "UNIQUE", "—", "Protect & lean in. Strongest emotional differentiator."),

 ("Echelon-Unique Strengths", "Game-world / spatial UIs",
  "Pallet Town overworld, 3D AgentWorld, infinite canvas, terminal grid.",
  "None.", "None.",
  "UNIQUE", "—", "Delightful, demo-able, unique. Polish the few state-machine bugs noted in codebase map."),

 ("Echelon-Unique Strengths", "Author-once, build-to-many targets",
  "Same team defs compile to Claude Code plugin, OpenClaw bundle, Echelon season-pack.",
  "Python-only.", "SaaS-only.",
  "UNIQUE", "—", "Portability/no-lock-in story competitors can't match. Finish openclaw pin (currently TBD)."),

 ("Echelon-Unique Strengths", "Advisory board (SME consult)",
  "12 tech-luminary advisors + Stephen Hawking oracle for high-level consults.",
  "None.", "None.",
  "UNIQUE", "—", "Distinctive. Wire it to the real counselor backend when that lands."),
]

# write rows with category grouping
row_idx = 2
prev_cat = None
n = 0
for cat, cap, ech, crew, hatz, tag, effort, note in ROWS:
    # category banner row when category changes
    if cat != prev_cat:
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=9)
        bc = ws.cell(row=row_idx, column=1, value=cat.upper())
        bc.fill = CAT_FILL; bc.font = CAT_FONT
        bc.alignment = Alignment(horizontal="left", vertical="center")
        bc.border = BORDER
        ws.row_dimensions[row_idx].height = 20
        row_idx += 1
        prev_cat = cat
    n += 1
    vals = [n, "", cap, ech, crew, hatz, tag, effort, note]
    for col, v in enumerate(vals, start=1):
        c = ws.cell(row=row_idx, column=col, value=v)
        c.alignment = CENTER if col in (1,7,8) else WRAP
        c.border = BORDER
        c.font = Font(size=10)
        if col == 7 and tag in TAG_FILL:
            c.fill = TAG_FILL[tag]; c.font = Font(size=9, bold=True)
        if col == 8 and effort in EFFORT_FILL:
            c.fill = EFFORT_FILL[effort]; c.font = Font(size=10, bold=True)
        if col == 3:
            c.font = Font(size=10, bold=True)
    ws.row_dimensions[row_idx].height = 58
    row_idx += 1

# ===========================================================================
# SHEET 3 — Build Roadmap (prioritised shortlist)
# ===========================================================================
ws = wb.create_sheet("Build Roadmap")
rcols = ["Priority", "Initiative", "Why (learning from CrewAI/Hatz)", "Effort", "Depends on"]
rwidths = [10, 34, 60, 9, 26]
for i,(c,w) in enumerate(zip(rcols,rwidths),start=1):
    cell = ws.cell(row=1,column=i,value=c)
    cell.fill=HEADER_FILL; cell.font=HEADER_FONT; cell.alignment=CENTER; cell.border=BORDER
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes="A2"
ROADMAP = [
 ("P0", "Make the Counselor real",
  "The 4-model consensus council is Echelon's #1 unique claim but returns MOCK verdicts. Wiring real parallel fan-out + aggregation turns marketing into product. No competitor has this.", "M", "Provider API keys (already wired from CLI auth)"),
 ("P0", "Implement the 7 review gates",
  "Gates are stubbed. Real gate runner (incl. a CrewAI-style hallucination/faithfulness scorer on the security gate) is the quality story that differentiates from both competitors.", "L", "Counselor (for escalation)"),
 ("P0", "Realize the knowledge flywheel",
  "Capture-on-merge + retrieve-before-task is the compounding-memory moat; partly stubbed (kb promote = v0.5). Borrow CrewAI's weighted/scoped memory ranking.", "L", "mempalace pin (placeholder SHA today)"),
 ("P1", "Deterministic workflow layer (Flows-equivalent)",
  "Typed, inspectable orchestration above seasons. Directly model on CrewAI Flows (@start/@listen/@router/@persist, Pydantic state). Makes runs debuggable & resumable.", "L", "—"),
 ("P1", "OpenTelemetry tracing + ops dashboard",
  "Biggest enterprise gap vs both. Structured spans for agent/tool/LLM calls unlock cost analytics, audit, and debugging. Foundation for everything enterprise.", "XL", "—"),
 ("P1", "Programmatic + LLM output guardrails & typed outputs",
  "CrewAI's function/LLM guardrails + Pydantic outputs are cheap, high-leverage reliability wins.", "M", "—"),
 ("P1", "Integration breadth: OAuth connector framework + Zapier/Make/n8n bridge",
  "Echelon has ~5 integrations vs 23-40+. The automation-bridge pattern (Hatz) reaches thousands cheaply.", "L", "—"),
 ("P2", "Agent-eval harness (LLM-scored)",
  "CrewAI 'crewai test' pattern: run a season N times, LLM-score per-character/task. Quantifies quality; feeds the flywheel.", "L", "Tracing"),
 ("P2", "Multi-format + audio knowledge ingestion",
  "Match Hatz's 60+ file types + audio and CrewAI's source classes + query rewriting.", "M", "KB flywheel"),
 ("P2", "Shareable season/blueprint registry",
  "Both competitors invest in templates/marketplaces for network effects. Echelon's YAML model is ready-made.", "L", "—"),
 ("P2", "Visual season/flow builder (no-code)",
  "Both competitors' biggest adoption lever. Broadens audience beyond developers.", "XL", "Workflow layer"),
 ("P3", "Append-only audit log + agent-capability→human RBAC bridge",
  "Pre-work for any team edition; useful for solo trust too.", "M", "—"),
 ("STRATEGIC", "Decide on a hosted/team edition",
  "Multi-tenancy, SSO/SCIM, SOC2, on-prem (CrewAI Factory / Hatz tenant model) ALL hinge on this one fork-in-the-road. Contradicts current local-first thesis — choose deliberately, don't drift.", "XXL", "Executive decision"),
]
ridx=2
PRI_FILL={"P0":PatternFill("solid",fgColor="FECACA"),"P1":PatternFill("solid",fgColor="FED7AA"),
          "P2":PatternFill("solid",fgColor="FEF3C7"),"P3":PatternFill("solid",fgColor="E0E7FF"),
          "STRATEGIC":PatternFill("solid",fgColor="EDE9FE")}
for pri,init,why,eff,dep in ROADMAP:
    vals=[pri,init,why,eff,dep]
    for col,v in enumerate(vals,start=1):
        c=ws.cell(row=ridx,column=col,value=v)
        c.alignment=CENTER if col in (1,4) else WRAP
        c.border=BORDER; c.font=Font(size=10)
        if col==1: c.fill=PRI_FILL.get(pri); c.font=Font(size=10,bold=True)
        if col==2: c.font=Font(size=10,bold=True)
        if col==4 and eff in EFFORT_FILL: c.fill=EFFORT_FILL[eff]; c.font=Font(size=10,bold=True)
    ws.row_dimensions[ridx].height=58
    ridx+=1

# ===========================================================================
# SHEET 4 — Platform Profiles
# ===========================================================================
ws = wb.create_sheet("Platform Profiles")
ws.sheet_view.showGridLines=False
pcols=["", "Echelon", "CrewAI", "Hatz AI"]
pw=[24,46,46,46]
for i,(c,w) in enumerate(zip(pcols,pw),start=1):
    cell=ws.cell(row=1,column=i,value=c)
    cell.fill=HEADER_FILL; cell.font=HEADER_FONT; cell.alignment=CENTER; cell.border=BORDER
    ws.column_dimensions[get_column_letter(i)].width=w
PROFILE=[
 ("What it is","Local-first Electron+Next.js desktop app hosting the factor-echelon team-factory.","Open-source Python multi-agent framework (MIT) + commercial AMP control plane.","Closed-source multi-tenant white-label SaaS sold through MSPs."),
 ("Primary user","Solo developer / small team running agents locally.","Python developers building agentic apps.","SMB business users, via MSP partners."),
 ("Delivery","Desktop app, BYO API keys.","pip framework + cloud/self-host AMP.","Managed SaaS (AWS US-East), white-labeled."),
 ("Orchestration","Themed seasons, archetypes, super-agent + convener, worktree isolation.","Crews (autonomous) + Flows (deterministic, event-driven).","No-code workflows, apps, agents, voice."),
 ("Memory/KB","mempalace KB (local), soul-package memory, knowledge flywheel (partly stubbed).","Unified weighted/scoped Memory + Knowledge/RAG (Chroma/Qdrant).","Memory RAG, Pinecone/Snowflake, 60+ file types + audio."),
 ("Models","Multi-CLI (Claude/Codex/Gemini/OpenCode/Pi) + local; 4-model Counselor (stubbed).","Native 6 providers + LiteLLM for the rest.","Mido gateway: 6 providers w/ cost multipliers."),
 ("Quality","7 review gates + 4-model counselor (both STUBBED in v0.1).","Hallucination guardrail, function/LLM guardrails, eval/train CLI.","Admin controls; no code-review pipeline."),
 ("Enterprise","None yet (single-user, local).","Control plane, RBAC, SSO, on-prem/VPC (Factory), HITL.","Multi-tenant, SSO+SCIM, RBAC, SOC2 I/II+III, admin portal."),
 ("Observability","Per-agent terminals, usage tracking; no tracing.","AMP traces (OTEL) + Langfuse/Phoenix/Weave integrations.","Audit trails, per-tenant usage dashboards."),
 ("Integrations","~5 (Telegram/Slack/GitHub/JIRA/X) + 7 MCP servers.","23+ OAuth + first-class MCP + Composio.","40+ incl. M365 deep, PSA/MSP, Zapier/Make/n8n, MCP."),
 ("Pricing","Free; ~$1-5 API cost/season; $50 counselor cap.","Free ~50 exec/mo; Pro ~$25; Enterprise custom.","Channel-gated, pooled non-rollover credits (~$190/mo entry)."),
 ("Compliance","Local-first (data never leaves machine); no telemetry.","Claims SOC2/HIPAA/FedRAMP (unverified publicly).","SOC2 Type I+II, SOC3; HIPAA/GDPR not public."),
 ("Funding/maturity","Pre-release v0.1.0; solo/indie.","~$18M (Insight Partners); large community.","~$5M seed; founded 2023; thin reference base."),
 ("Signature moat","Themed characters, multi-model counselor, multi-target build, local-first.","Crews/Flows split, huge community, blueprints.","MSP channel, true white-label multi-tenancy."),
 ("Key weakness","Flagship features stubbed; no enterprise/observability/integrations breadth.","Hierarchical reliability, token cost, weak native observability, 2026 CVEs.","No public pricing, no 3rd-party reviews, GDPR/HIPAA unconfirmed."),
]
pr=2
for label,e,c,h in PROFILE:
    cells=[label,e,c,h]
    for col,v in enumerate(cells,start=1):
        cell=ws.cell(row=pr,column=col,value=v)
        cell.alignment=WRAP; cell.border=BORDER; cell.font=Font(size=10)
        if col==1: cell.font=Font(size=10,bold=True); cell.fill=PatternFill("solid",fgColor="F3F4F6")
    ws.row_dimensions[pr].height=46
    pr+=1

import os
out_dir="docs/research"
os.makedirs(out_dir,exist_ok=True)
out=os.path.join(out_dir,"2026-06-15-echelon-crewai-hatz-comparison.xlsx")
wb.save(out)
print("wrote",out)
print("sheets:",wb.sheetnames)
print("feature rows:",n)

# also emit a CSV of the feature matrix for portability / diffing
import csv
csv_out=os.path.join(out_dir,"2026-06-15-echelon-crewai-hatz-comparison.csv")
with open(csv_out,"w",newline="") as f:
    w=csv.writer(f)
    w.writerow(["#","Category","Capability","Echelon (v0.1)","CrewAI","Hatz AI","Alignment","Effort","Recommendation / Notes"])
    for i,(cat,cap,ech,crew,hatz,tag,effort,note) in enumerate(ROWS,start=1):
        w.writerow([i,cat,cap,ech,crew,hatz,tag,effort,note])
print("wrote",csv_out)
